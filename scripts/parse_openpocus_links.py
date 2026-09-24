"""
Parse 'Link to US Images.docx' from kumarandre/OpenPOCUS to extract image URLs.

Usage:
  python scripts/parse_openpocus_links.py
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    sys.exit("Run: pip install python-docx")

DOCX_PATH = Path("data/raw/openpocus_repo/Link to US Images.docx")


def extract_links(docx_path: Path) -> list:
    doc = Document(str(docx_path))
    links = []

    for para in doc.paragraphs:
        urls = re.findall(r'https?://\S+', para.text)
        links.extend(urls)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                urls = re.findall(r'https?://\S+', cell.text)
                links.extend(urls)

    for rel in doc.part.rels.values():
        if "hyperlink" in rel.reltype:
            target = rel._target
            if target.startswith("http"):
                links.append(target)

    seen = set()
    unique = []
    for l in links:
        l = l.rstrip(".,);")
        if l not in seen:
            seen.add(l)
            unique.append(l)

    return unique


def main():
    if not DOCX_PATH.exists():
        sys.exit(f"Not found: {DOCX_PATH}")

    links = extract_links(DOCX_PATH)
    print(f"Found {len(links)} links:")
    for l in links:
        print(l)

    print("\n--- Full document text ---")
    doc = Document(str(DOCX_PATH))
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            print(f"[{i}] {para.text}")
    for table in doc.tables:
        print("\n--- Table ---")
        for row in table.rows:
            print(" | ".join(cell.text.strip() for cell in row.cells))


if __name__ == "__main__":
    main()
